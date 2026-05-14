from pathlib import Path
import json
import sqlite3
from io import BytesIO
from urllib.parse import urlparse

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import history as history_api_module
from src.api import procurement as procurement_api_module
from src.api.main import app
from src.api.routers import copilot as copilot_router_module
from src.services.copilot_context import build_analysis_assist_context, build_import_assist_context
from src.services.ollama_client import OllamaClientError

client = TestClient(app)


@pytest.fixture(autouse=True)
def isolate_history_storage(tmp_path: Path):
    history_api_module.repository.db_path = tmp_path / "history.db"
    history_api_module.repository.legacy_json_path = tmp_path / "history.json"
    history_api_module.repository._migration_completed = False
    history_api_module.ensure_history_file()
    procurement_api_module.repository.db_path = tmp_path / "history.db"
    procurement_api_module.ensure_procurement_storage()
    yield


def _create_test_sales_file(tmp_path: Path) -> Path:
    input_file = tmp_path / "api_sales.csv"

    df = pd.DataFrame(
        [
            {
                "date": "2026-01-01",
                "sku": "COF-001",
                "product_name": "Кофе арабика 250г",
                "category": "Кофе",
                "quantity_sold": 5,
                "revenue": 2100,
                "stock_on_hand": 120,
                "purchase_price": 250,
                "selling_price": 420,
            },
            {
                "date": "2026-01-02",
                "sku": "COF-001",
                "product_name": "Кофе арабика 250г",
                "category": "Кофе",
                "quantity_sold": 6,
                "revenue": 2520,
                "stock_on_hand": 114,
                "purchase_price": 250,
                "selling_price": 420,
            },
            {
                "date": "2026-01-01",
                "sku": "TEA-001",
                "product_name": "Чай зелёный 100г",
                "category": "Чай",
                "quantity_sold": 3,
                "revenue": 990,
                "stock_on_hand": 80,
                "purchase_price": 160,
                "selling_price": 330,
            },
            {
                "date": "2026-01-02",
                "sku": "TEA-001",
                "product_name": "Чай зелёный 100г",
                "category": "Чай",
                "quantity_sold": 4,
                "revenue": 1320,
                "stock_on_hand": 76,
                "purchase_price": 160,
                "selling_price": 330,
            },
        ]
    )

    df.to_csv(input_file, index=False)

    return input_file


def _create_missing_stock_file(tmp_path: Path) -> Path:
    input_file = tmp_path / "missing_stock.csv"
    df = pd.DataFrame(
        [
            {
                "date": "2026-01-01",
                "sku": "COF-001",
                "product_name": "Кофе арабика 250г",
                "quantity_sold": 5,
                "revenue": 2100,
            }
        ]
    )
    df.to_csv(input_file, index=False)
    return input_file


def _create_test_sales_xlsx(tmp_path: Path) -> Path:
    input_file = tmp_path / "api_sales.xlsx"
    df = pd.read_csv(_create_test_sales_file(tmp_path))
    df.to_excel(input_file, index=False)
    return input_file


def _create_ru_named_sales_file(tmp_path: Path) -> Path:
    input_file = tmp_path / "sales_ru.csv"
    df = pd.DataFrame(
        [
            {
                "Дата продажи": "2026-01-01",
                "Артикул": "COF-001",
                "Товар": "Кофе арабика 250г",
                "Продано": 5,
                "Сумма": 2100,
                "Остаток": 120,
            },
            {
                "Дата продажи": "2026-01-02",
                "Артикул": "TEA-001",
                "Товар": "Чай зелёный 100г",
                "Продано": 4,
                "Сумма": 1320,
                "Остаток": 76,
            },
        ]
    )
    df.to_csv(input_file, index=False)
    return input_file


def _create_en_named_sales_file(tmp_path: Path) -> Path:
    input_file = tmp_path / "sales_en.csv"
    df = pd.DataFrame(
        [
            {
                "sale_date": "2026-01-01",
                "item_code": "COF-001",
                "product": "Кофе арабика 250г",
                "units_sold": 5,
                "amount": 2100,
                "inventory": 120,
            }
        ]
    )
    df.to_csv(input_file, index=False)
    return input_file


def test_health_endpoint():
    response = client.get("/health")

    assert response.status_code == 200

    data = response.json()

    assert data["status"] == "ok"
    assert data["service"] == "StockPilot AI API"


def test_validate_file_endpoint_success(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/validate-file",
            files={"file": ("api_sales.csv", file, "text/csv")},
        )

    assert response.status_code == 200

    data = response.json()

    assert data["status"] == "success"
    assert data["is_valid"] is True
    assert data["row_count"] == 4
    assert data["missing_columns"] == []
    assert data["errors"] == []
    assert "data_quality" in data
    assert data["data_quality"]["total_rows"] == 4


def test_import_inspect_endpoint_success_for_valid_csv(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/import/inspect",
            files={"file": ("api_sales.csv", file, "text/csv")},
        )

    assert response.status_code == 200

    data = response.json()

    assert data["filename"] == "api_sales.csv"
    assert data["file_type"] == "csv"
    assert data["row_count"] == 4
    assert data["column_count"] >= 6
    assert data["is_ready_for_analysis"] is True
    assert data["required_columns_status"]["date"]["present"] is True
    assert data["required_columns_status"]["stock_on_hand"]["present"] is True
    assert len(data["sample_rows"]) == 4
    assert data["sku_count"] == 2
    assert data["date_range"]["min"] == "2026-01-01"
    assert data["date_range"]["max"] == "2026-01-02"


def test_import_inspect_endpoint_detects_missing_required_column(tmp_path: Path):
    input_file = _create_missing_stock_file(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/import/inspect",
            files={"file": ("missing_stock.csv", file, "text/csv")},
        )

    assert response.status_code == 200

    data = response.json()

    assert data["is_ready_for_analysis"] is False
    assert data["required_columns_status"]["stock_on_hand"]["present"] is False
    assert any(
        warning["code"] == "missing_required_column" and warning["field"] == "stock_on_hand"
        for warning in data["warnings"]
    )


def test_import_inspect_endpoint_rejects_unsupported_file_type(tmp_path: Path):
    input_file = tmp_path / "notes.txt"
    input_file.write_text("hello", encoding="utf-8")

    with input_file.open("rb") as file:
        response = client.post(
            "/import/inspect",
            files={"file": ("notes.txt", file, "text/plain")},
        )

    assert response.status_code == 400


def test_import_inspect_endpoint_supports_xlsx(tmp_path: Path):
    input_file = _create_test_sales_xlsx(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/import/inspect",
            files={
                "file": (
                    "api_sales.xlsx",
                    file,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200

    data = response.json()

    assert data["file_type"] == "xlsx"
    assert data["is_ready_for_analysis"] is True
    assert data["required_columns_status"]["sku"]["present"] is True


def test_import_suggest_mapping_matches_russian_columns(tmp_path: Path):
    input_file = _create_ru_named_sales_file(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/import/suggest-mapping",
            files={"file": ("sales_ru.csv", file, "text/csv")},
        )

    assert response.status_code == 200
    data = response.json()

    assert data["suggested_mapping"]["date"]["source_column"] == "Дата продажи"
    assert data["suggested_mapping"]["sku"]["source_column"] == "Артикул"
    assert data["suggested_mapping"]["product_name"]["source_column"] == "Товар"
    assert data["suggested_mapping"]["quantity_sold"]["source_column"] == "Продано"
    assert data["suggested_mapping"]["revenue"]["source_column"] == "Сумма"
    assert data["suggested_mapping"]["stock_on_hand"]["source_column"] == "Остаток"
    assert data["is_ready_for_normalization"] is True


def test_import_suggest_mapping_matches_english_synonyms(tmp_path: Path):
    input_file = _create_en_named_sales_file(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/import/suggest-mapping",
            files={"file": ("sales_en.csv", file, "text/csv")},
        )

    assert response.status_code == 200
    data = response.json()

    assert data["suggested_mapping"]["date"]["source_column"] == "sale_date"
    assert data["suggested_mapping"]["sku"]["source_column"] == "item_code"
    assert data["suggested_mapping"]["product_name"]["source_column"] == "product"
    assert data["suggested_mapping"]["quantity_sold"]["source_column"] == "units_sold"
    assert data["suggested_mapping"]["revenue"]["source_column"] == "amount"
    assert data["suggested_mapping"]["stock_on_hand"]["source_column"] == "inventory"


def test_import_suggest_mapping_reports_missing_required_fields(tmp_path: Path):
    input_file = _create_missing_stock_file(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/import/suggest-mapping",
            files={"file": ("missing_stock.csv", file, "text/csv")},
        )

    assert response.status_code == 200
    data = response.json()

    assert data["is_ready_for_normalization"] is False
    assert data["suggested_mapping"]["stock_on_hand"]["source_column"] is None


def test_import_validate_mapping_valid_case(tmp_path: Path):
    input_file = _create_ru_named_sales_file(tmp_path)
    mapping = {
        "date": {"source_column": "Дата продажи"},
        "sku": {"source_column": "Артикул"},
        "product_name": {"source_column": "Товар"},
        "quantity_sold": {"source_column": "Продано"},
        "revenue": {"source_column": "Сумма"},
        "stock_on_hand": {"source_column": "Остаток"},
    }

    with input_file.open("rb") as file:
        response = client.post(
            "/import/validate-mapping",
            files={"file": ("sales_ru.csv", file, "text/csv")},
            data={"mapping": json.dumps(mapping, ensure_ascii=False)},
        )

    assert response.status_code == 200
    data = response.json()
    assert data["is_valid"] is True
    assert data["errors"] == []


def test_import_validate_mapping_rejects_unknown_source_column(tmp_path: Path):
    input_file = _create_ru_named_sales_file(tmp_path)
    mapping = {
        "date": {"source_column": "Дата продажи"},
        "sku": {"source_column": "Несуществующая колонка"},
        "product_name": {"source_column": "Товар"},
        "quantity_sold": {"source_column": "Продано"},
        "revenue": {"source_column": "Сумма"},
        "stock_on_hand": {"source_column": "Остаток"},
    }

    with input_file.open("rb") as file:
        response = client.post(
            "/import/validate-mapping",
            files={"file": ("sales_ru.csv", file, "text/csv")},
            data={"mapping": json.dumps(mapping, ensure_ascii=False)},
        )

    assert response.status_code == 200
    data = response.json()
    assert data["is_valid"] is False
    assert any(error["code"] == "missing_source_column" for error in data["errors"])


def test_import_normalize_csv_and_analyze_roundtrip(tmp_path: Path):
    input_file = _create_ru_named_sales_file(tmp_path)
    mapping = {
        "date": {"source_column": "Дата продажи"},
        "sku": {"source_column": "Артикул"},
        "product_name": {"source_column": "Товар"},
        "quantity_sold": {"source_column": "Продано"},
        "revenue": {"source_column": "Сумма"},
        "stock_on_hand": {"source_column": "Остаток"},
    }

    with input_file.open("rb") as file:
        normalize_response = client.post(
            "/import/normalize",
            files={"file": ("sales_ru.csv", file, "text/csv")},
            data={"mapping": json.dumps(mapping, ensure_ascii=False)},
        )

    assert normalize_response.status_code == 200
    normalized_text = normalize_response.content.decode("utf-8-sig")
    assert "date,sku,product_name,quantity_sold,revenue,stock_on_hand" in normalized_text
    assert normalized_text.count("\n") >= 2

    analyze_response = client.post(
        "/analyze",
        files={"file": ("normalized.csv", BytesIO(normalize_response.content), "text/csv")},
        data={"horizon": "14", "safety_stock": "0.2", "lead_time": "3"},
    )

    assert analyze_response.status_code == 200


def test_import_wizard_rejects_unsupported_file_type(tmp_path: Path):
    input_file = tmp_path / "notes.txt"
    input_file.write_text("hello", encoding="utf-8")

    with input_file.open("rb") as file:
        response = client.post(
            "/import/suggest-mapping",
            files={"file": ("notes.txt", file, "text/plain")},
        )

    assert response.status_code == 400


def test_import_assist_endpoint_returns_mocked_ollama_response(monkeypatch: pytest.MonkeyPatch):
    async def fake_generate_import_assist_answer(question: str, context: dict[str, object]) -> tuple[str, str]:
        assert question == "Что не так с моим файлом?"
        assert context["filename"] == "sales_ru.csv"
        return ("Нужно проверить колонку «Остаток».", "qwen2.5:7b")

    monkeypatch.setattr(
        copilot_router_module,
        "generate_import_assist_answer",
        fake_generate_import_assist_answer,
    )

    response = client.post(
        "/copilot/import-assist",
        json={
            "question": "Что не так с моим файлом?",
            "context": {
                "stage": "upload_mapping",
                "filename": "sales_ru.csv",
                "columns": ["Дата продажи", "Артикул", "Товар", "Продано", "Сумма"],
                "required_columns_status": {"date": {"present": False, "source_column": None}},
                "warnings": [],
                "sample_rows": [{"Дата продажи": "2026-01-01", "Артикул": "COF-001"}],
            },
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["answer"] == "Нужно проверить колонку «Остаток»."
    assert data["model"] == "qwen2.5:7b"
    assert data["provider"] == "ollama"
    assert data["is_local"] is True
    assert data["fallback"] is False


def test_import_assist_endpoint_returns_fallback_when_ollama_unavailable(monkeypatch: pytest.MonkeyPatch):
    async def fake_generate_import_assist_answer(question: str, context: dict[str, object]) -> tuple[str, str]:
        raise OllamaClientError("unavailable")

    monkeypatch.setattr(
        copilot_router_module,
        "generate_import_assist_answer",
        fake_generate_import_assist_answer,
    )

    response = client.post(
        "/copilot/import-assist",
        json={
            "question": "Почему файл не готов к анализу?",
            "context": {"filename": "sales_ru.csv"},
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["fallback"] is True
    assert "Локальный AI-помощник сейчас недоступен" in data["answer"]
    assert data["provider"] == "ollama"
    assert data["error_code"] == "unavailable"


def test_import_assist_does_not_require_openai_config(monkeypatch: pytest.MonkeyPatch):
    async def fake_generate_import_assist_answer(question: str, context: dict[str, object]) -> tuple[str, str]:
        return ("Контекст получен локально.", "qwen2.5:7b")

    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(
        copilot_router_module,
        "generate_import_assist_answer",
        fake_generate_import_assist_answer,
    )

    response = client.post(
        "/copilot/import-assist",
        json={
            "question": "Какие колонки обязательны?",
            "context": {"filename": "sales.csv", "columns": ["date", "sku"]},
        },
    )

    assert response.status_code == 200
    assert response.json()["fallback"] is False


def test_import_assist_context_builder_keeps_context_compact():
    context = build_import_assist_context(
        {
            "stage": "upload_mapping",
            "filename": "sales_ru.csv",
            "columns": ["Дата продажи", "Артикул", "Товар"],
            "required_columns_status": {"date": {"present": True, "source_column": "Дата продажи"}},
            "optional_columns_status": {"supplier": {"present": False, "source_column": None}},
            "suggested_mapping": {"date": {"source_column": "Дата продажи", "confidence": 0.98, "status": "matched"}},
            "validation_errors": [{"code": "missing", "title": "Ошибка", "message": "Не хватает поля"}],
            "warnings": [{"code": "warn", "title": "Внимание", "message": "Проверьте файл"}],
            "sample_rows": [
                {"Дата продажи": "2026-01-01", "Артикул": "COF-001"},
                {"Дата продажи": "2026-01-02", "Артикул": "TEA-001"},
                {"Дата продажи": "2026-01-03", "Артикул": "MIL-001"},
                {"Дата продажи": "2026-01-04", "Артикул": "SNK-001"},
                {"Дата продажи": "2026-01-05", "Артикул": "JCE-001"},
                {"Дата продажи": "2026-01-06", "Артикул": "WTR-001"},
            ],
            "full_dataset": [{"very": "large"}],
            "secret": "must_not_pass",
        }
    )

    assert "full_dataset" not in context
    assert "secret" not in context
    assert len(context["sample_rows"]) == 5
    assert context["filename"] == "sales_ru.csv"


def test_import_assist_system_prompt_has_local_only_rules():
    prompt = copilot_router_module.build_import_assist_system_prompt()

    assert "Ты отвечаешь только на основе переданного контекста" in prompt
    assert "Ты не выдумываешь колонки, значения, метрики или выводы" in prompt
    assert "Не упоминай OpenAI" in prompt
    assert "Не предлагай внешние AI-сервисы" in prompt


def test_analysis_assist_endpoint_returns_mocked_ollama_response(monkeypatch: pytest.MonkeyPatch):
    async def fake_generate_analysis_assist_answer(question: str, context: dict[str, object]) -> tuple[str, str]:
        assert question == "Что нужно сделать в первую очередь?"
        assert context["stage"] == "analysis"
        return ("Сначала проверьте SKU с высоким риском дефицита.", "qwen2.5:7b")

    monkeypatch.setattr(
        copilot_router_module,
        "generate_analysis_assist_answer",
        fake_generate_analysis_assist_answer,
    )

    response = client.post(
        "/copilot/analysis-assist",
        json={
            "question": "Что нужно сделать в первую очередь?",
            "context": {
                "stage": "analysis",
                "summary": {"filename": "sales.csv", "period": "2026-01-01 - 2026-01-31", "sku_count": 12, "row_count": 400},
                "top_recommendations": [{"sku": "COF-001", "recommended_order_qty": 10}],
            },
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["answer"] == "Сначала проверьте SKU с высоким риском дефицита."
    assert data["fallback"] is False


def test_analysis_assist_endpoint_returns_fallback_when_ollama_unavailable(monkeypatch: pytest.MonkeyPatch):
    async def fake_generate_analysis_assist_answer(question: str, context: dict[str, object]) -> tuple[str, str]:
        raise OllamaClientError("timeout")

    monkeypatch.setattr(
        copilot_router_module,
        "generate_analysis_assist_answer",
        fake_generate_analysis_assist_answer,
    )

    response = client.post(
        "/copilot/analysis-assist",
        json={"question": "Какие товары срочно докупить?", "context": {"stage": "analysis"}},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["fallback"] is True
    assert "не ответил вовремя" in data["answer"]
    assert data["error_code"] == "timeout"


def test_copilot_status_endpoint_returns_available_status(monkeypatch: pytest.MonkeyPatch):
    async def fake_get_status(self):
        return (True, "Ollama доступна, модель qwen2.5:7b готова.", None)

    monkeypatch.setattr("src.services.ollama_client.OllamaClient.get_status", fake_get_status)

    response = client.get("/copilot/status")

    assert response.status_code == 200
    data = response.json()
    assert data["provider"] == "ollama"
    assert data["is_local"] is True
    assert data["available"] is True
    assert data["model"] == "qwen2.5:7b"
    assert data["error_code"] is None


def test_copilot_status_endpoint_returns_missing_model_state(monkeypatch: pytest.MonkeyPatch):
    async def fake_get_status(self):
        return (False, "Модель qwen2.5:7b не найдена в локальной Ollama.", "missing_model")

    monkeypatch.setattr("src.services.ollama_client.OllamaClient.get_status", fake_get_status)

    response = client.get("/copilot/status")

    assert response.status_code == 200
    data = response.json()
    assert data["available"] is False
    assert data["error_code"] == "missing_model"
    assert "не найдена" in data["message"]


def test_copilot_status_endpoint_handles_unavailable_ollama(monkeypatch: pytest.MonkeyPatch):
    async def fake_get_status(self):
        raise OllamaClientError("unavailable")

    monkeypatch.setattr("src.services.ollama_client.OllamaClient.get_status", fake_get_status)

    response = client.get("/copilot/status")

    assert response.status_code == 200
    data = response.json()
    assert data["available"] is False
    assert data["error_code"] == "unavailable"
    assert "Ollama запущена" in data["message"]


def test_analysis_assist_does_not_require_openai_config(monkeypatch: pytest.MonkeyPatch):
    async def fake_generate_analysis_assist_answer(question: str, context: dict[str, object]) -> tuple[str, str]:
        return ("Контекст анализа получен локально.", "qwen2.5:7b")

    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(
        copilot_router_module,
        "generate_analysis_assist_answer",
        fake_generate_analysis_assist_answer,
    )

    response = client.post(
        "/copilot/analysis-assist",
        json={"question": "Что происходит с рисками?", "context": {"stage": "analysis"}},
    )

    assert response.status_code == 200
    assert response.json()["fallback"] is False


def test_analysis_assist_context_builder_limits_arrays_and_keeps_warnings():
    context = build_analysis_assist_context(
        {
            "stage": "analysis",
            "summary": {"filename": "sales.csv", "period": "2026-01", "sku_count": 50, "row_count": 5000},
            "top_recommendations": [{"sku": f"SKU-{index:03d}"} for index in range(30)],
            "high_risk_items": [{"sku": f"HR-{index:03d}"} for index in range(30)],
            "reliability_warnings": [
                {
                    "code": "possible_stockout_censored_demand",
                    "title": "Возможное искажение спроса",
                    "message": "Проверьте дни с нулевым остатком",
                }
            ],
            "raw_dataset": [{"huge": "table"}],
            "selected_sku": {"sku": "COF-001", "confidence": "low"},
        }
    )

    assert len(context["top_recommendations"]) == 15
    assert len(context["high_risk_items"]) == 15
    assert "raw_dataset" not in context
    assert context["reliability_warnings"][0]["code"] == "possible_stockout_censored_demand"
    assert context["selected_sku"]["sku"] == "COF-001"


def test_analysis_assist_system_prompt_has_local_only_rules():
    prompt = copilot_router_module.build_analysis_assist_system_prompt()

    assert "Отвечай только на основе переданного контекста" in prompt
    assert "Не выдумывай SKU, цифры, метрики, поставщиков, прогнозы или рекомендации" in prompt
    assert "Если forecast confidence низкий, обязательно упомяни это" in prompt
    assert "Не упоминай OpenAI" in prompt


def test_analyze_endpoint_returns_reports_and_csv(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={
                "horizon": "14",
                "safety_stock": "0.2",
                "lead_time": "3",
            },
        )

    assert response.status_code == 200

    data = response.json()

    assert data["status"] == "success"
    assert data["summary_metrics"]["total_revenue"] == 6930
    assert len(data["recommendations"]) > 0
    assert len(data["top_products"]) > 0
    assert len(data["product_forecast"]) > 0
    assert "analysis_reliability" in data
    assert "forecast_quality" in data["analysis_reliability"]
    assert "sku_reliability" in data["analysis_reliability"]
    assert len(data["abc_analysis"]) > 0
    assert "abc_summary" in data
    assert len(data["xyz_analysis"]) > 0
    assert "xyz_summary" in data
    assert len(data["abc_xyz_matrix"]) > 0
    assert "abc_xyz_summary" in data

    assert data["report_url"].endswith(".xlsx")
    assert data["dashboard_url"].endswith(".html")
    assert data["recommendations_csv_url"].endswith(".csv")
    assert data["abc_xyz_matrix_csv_url"].endswith(".csv")

    report_path = Path(data["report_path"])

    assert report_path.exists()
    assert Path(data["dashboard_path"]).exists()
    assert Path(data["recommendations_csv_path"]).exists()
    assert Path(data["abc_xyz_matrix_csv_path"]).exists()

    workbook = load_workbook(report_path)

    assert "Сводка" in workbook.sheetnames
    assert "Проблемы данных" in workbook.sheetnames
    assert "Надёжность анализа" in workbook.sheetnames
    assert "Закупочный план" in workbook.sheetnames
    assert "Технические рекомендации" in workbook.sheetnames
    assert "ABC анализ" in workbook.sheetnames
    assert "XYZ анализ" in workbook.sheetnames
    assert "ABC XYZ матрица" in workbook.sheetnames
    assert len(workbook["Сводка"]._charts) >= 3

    dashboard_html = Path(data["dashboard_path"]).read_text(encoding="utf-8")
    assert "Управленческая сводка" in dashboard_html
    assert "Продажи и выручка" in dashboard_html
    assert "Закупочный план" in dashboard_html
    assert "Надёжность прогноза" in dashboard_html
    assert "Качество данных" in dashboard_html
    assert "ABC/XYZ" in dashboard_html
    assert "<svg" in dashboard_html
    assert "Executive Summary" not in dashboard_html
    assert "Sales & Revenue" not in dashboard_html
    assert "Procurement Plan" not in dashboard_html
    assert "Forecast Reliability" not in dashboard_html
    assert "Data Quality" not in dashboard_html
    assert "Ollama" not in dashboard_html
    assert "no external AI services" not in dashboard_html
    assert "local-first · Ollama" not in dashboard_html
    assert "Medium приоритет" not in dashboard_html
    assert "exp_smoothing" not in dashboard_html
    assert "moving_average_7" not in dashboard_html
    assert " stable" not in dashboard_html
    assert "269220.0" not in dashboard_html
    assert ">medium<" not in dashboard_html
    assert ">stable<" not in dashboard_html


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("horizon", "0"),
        ("horizon", "366"),
        ("safety_stock", "-0.1"),
        ("lead_time", "-1"),
    ],
)
def test_analyze_endpoint_rejects_invalid_parameters(
    tmp_path: Path,
    field: str,
    value: str,
):
    input_file = _create_test_sales_file(tmp_path)
    data = {
        "horizon": "14",
        "safety_stock": "0.2",
        "lead_time": "3",
        field: value,
    }

    with input_file.open("rb") as file:
        response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data=data,
        )

    assert response.status_code == 422


def test_download_recommendations_csv_endpoint(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        analyze_response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={
                "horizon": "14",
                "safety_stock": "0.2",
                "lead_time": "3",
            },
        )

    assert analyze_response.status_code == 200

    analyze_data = analyze_response.json()
    csv_url = analyze_data["recommendations_csv_url"]
    csv_path = urlparse(csv_url).path

    download_response = client.get(csv_path)

    assert download_response.status_code == 200
    assert "text/csv" in download_response.headers["content-type"]
    assert "SKU,Товар,Категория" in download_response.text
    assert "Рекомендовано системой" in download_response.text
    assert "Итоговое количество" in download_response.text
    assert "Приоритет" in download_response.text
    assert "Статус" in download_response.text
    assert "Новая" in download_response.text
    assert "product_name" not in download_response.text
    assert "Medium приоритет" not in download_response.text


def test_download_abc_xyz_matrix_csv_endpoint(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        analyze_response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={
                "horizon": "14",
                "safety_stock": "0.2",
                "lead_time": "3",
            },
        )

    assert analyze_response.status_code == 200

    analyze_data = analyze_response.json()
    csv_url = analyze_data["abc_xyz_matrix_csv_url"]
    csv_path = urlparse(csv_url).path

    download_response = client.get(csv_path)

    assert download_response.status_code == 200
    assert "text/csv" in download_response.headers["content-type"]
    assert "Класс ABC/XYZ" in download_response.text
    assert "Действие" in download_response.text


def test_download_recommendations_raw_csv_endpoint(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        analyze_response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={
                "horizon": "14",
                "safety_stock": "0.2",
                "lead_time": "3",
            },
        )

    assert analyze_response.status_code == 200

    analyze_data = analyze_response.json()
    csv_url = analyze_data["recommendations_raw_csv_url"]
    csv_path = urlparse(csv_url).path

    download_response = client.get(csv_path)

    assert download_response.status_code == 200
    assert "text/csv" in download_response.headers["content-type"]
    assert "product_name" in download_response.text
    assert "recommended_order_qty" in download_response.text


def test_abc_xyz_csv_contains_all_skus(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        analyze_response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={
                "horizon": "14",
                "safety_stock": "0.2",
                "lead_time": "3",
            },
        )

    assert analyze_response.status_code == 200
    analyze_data = analyze_response.json()

    matrix_path = Path(analyze_data["abc_xyz_matrix_csv_path"])
    matrix_df = pd.read_csv(matrix_path)

    assert len(matrix_df) == analyze_data["summary_metrics"]["unique_products"]


def test_history_endpoint_returns_pagination_metadata(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        analyze_response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={
                "horizon": "14",
                "safety_stock": "0.2",
                "lead_time": "3",
            },
        )

    assert analyze_response.status_code == 200

    history_response = client.get("/history?page=1&page_size=5")

    assert history_response.status_code == 200

    history_data = history_response.json()

    assert history_data["status"] == "success"
    assert "pagination" in history_data
    assert history_data["pagination"]["page"] == 1
    assert history_data["pagination"]["page_size"] == 5
    assert len(history_data["records"]) >= 1


def test_history_detail_and_result_endpoints_return_saved_snapshot(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        analyze_response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={"horizon": "14", "safety_stock": "0.2", "lead_time": "3"},
        )

    assert analyze_response.status_code == 200
    analyze_data = analyze_response.json()
    run_id = analyze_data["run_id"]

    detail_response = client.get(f"/history/{run_id}")
    assert detail_response.status_code == 200
    detail_data = detail_response.json()
    assert detail_data["id"] == run_id
    assert detail_data["original_filename"] == "api_sales.csv"
    assert detail_data["has_result_snapshot"] is True

    result_response = client.get(f"/history/{run_id}/result")
    assert result_response.status_code == 200
    result_data = result_response.json()
    assert result_data["run_id"] == run_id
    assert result_data["original_filename"] == "api_sales.csv"
    assert "recommendations" in result_data
    assert "analysis_reliability" in result_data


def test_history_files_and_decisions_endpoints(tmp_path: Path):
    input_file = _create_test_sales_file(tmp_path)

    with input_file.open("rb") as file:
        analyze_response = client.post(
            "/analyze",
            files={"file": ("api_sales.csv", file, "text/csv")},
            data={"horizon": "14", "safety_stock": "0.2", "lead_time": "3"},
        )

    run_id = analyze_response.json()["run_id"]

    files_response = client.get(f"/history/{run_id}/files")
    assert files_response.status_code == 200
    files_data = files_response.json()
    assert files_data["run_id"] == run_id
    assert len(files_data["files"]) == 4
    assert any(item["kind"] == "report" and item["available"] for item in files_data["files"])

    decisions_payload = {
        "COF-001": {
            "status": "accepted",
            "finalOrderQty": 18,
            "finalOrderQtyInput": "18",
            "comment": "Подтвердить у поставщика",
        }
    }
    save_response = client.put(f"/history/{run_id}/decisions", json=decisions_payload)
    assert save_response.status_code == 200
    assert save_response.json()["decisions"]["COF-001"]["status"] == "accepted"

    load_response = client.get(f"/history/{run_id}/decisions")
    assert load_response.status_code == 200
    assert load_response.json()["decisions"]["COF-001"]["comment"] == "Подтвердить у поставщика"


def test_history_endpoints_handle_broken_snapshot_and_missing_run(tmp_path: Path):
    history_api_module.repository.add_record(
        {
            "id": "broken-run",
            "original_filename": "broken.csv",
            "horizon": 14,
            "safety_stock": 0.2,
            "lead_time": 3,
            "summary_metrics": {"unique_products": 1},
            "recommendations_count": 0,
            "report_path": "",
            "dashboard_path": "",
            "report_url": "",
            "dashboard_url": "",
            "recommendations_csv_path": "",
            "recommendations_csv_url": "",
            "abc_xyz_matrix_csv_path": "",
            "abc_xyz_matrix_csv_url": "",
            "result_snapshot": {"status": "success", "message": "ok"},
        }
    )

    db_path = history_api_module.repository.db_path
    with sqlite3.connect(db_path) as connection:
        connection.execute(
            "UPDATE analysis_history SET result_snapshot_json = ?, summary_metrics_json = ? WHERE id = ?",
            ("{broken json", "{broken json", "broken-run"),
        )
        connection.commit()

    history_response = client.get("/history")
    assert history_response.status_code == 200
    assert history_response.json()["records"][0]["id"] == "broken-run"

    result_response = client.get("/history/broken-run/result")
    assert result_response.status_code == 404
    assert "сохранённый результат" in result_response.json()["detail"]["message"]

    missing_response = client.get("/history/missing-run")
    assert missing_response.status_code == 404
    assert "run_id" in missing_response.json()["detail"]["message"]

    missing_decisions_response = client.put("/history/missing-run/decisions", json={})
    assert missing_decisions_response.status_code == 404


def test_suppliers_crud_and_validation():
    invalid_response = client.post(
        "/suppliers",
        json={"name": "", "default_lead_time_days": -1},
    )
    assert invalid_response.status_code == 422

    create_response = client.post(
        "/suppliers",
        json={
            "name": "Поставщик А",
            "contact_name": "Ирина",
            "email": "supply@example.com",
            "default_lead_time_days": 7,
            "min_order_value": 10000,
            "notes": "Основной поставщик",
        },
    )
    assert create_response.status_code == 200
    supplier = create_response.json()
    assert supplier["name"] == "Поставщик А"
    supplier_id = supplier["id"]

    list_response = client.get("/suppliers")
    assert list_response.status_code == 200
    assert len(list_response.json()["suppliers"]) == 1

    update_response = client.put(
        f"/suppliers/{supplier_id}",
        json={
            "name": "Поставщик A+",
            "default_lead_time_days": 5,
        },
    )
    assert update_response.status_code == 200
    assert update_response.json()["name"] == "Поставщик A+"

    delete_response = client.delete(f"/suppliers/{supplier_id}")
    assert delete_response.status_code == 200
    assert client.get("/suppliers").json()["suppliers"] == []


def test_sku_settings_crud_and_validation():
    supplier_response = client.post(
        "/suppliers",
        json={"name": "Поставщик SKU", "default_lead_time_days": 4},
    )
    supplier_id = supplier_response.json()["id"]

    invalid_response = client.put(
        "/sku-settings/COF-001",
        json={"pack_size": 0, "supplier_id": "missing"},
    )
    assert invalid_response.status_code == 422

    create_response = client.put(
        "/sku-settings/COF-001",
        json={
            "product_name": "Кофе",
            "supplier_id": supplier_id,
            "purchase_price": 250,
            "selling_price": 420,
            "lead_time_days": 3,
            "moq": 30,
            "pack_size": 12,
            "safety_stock_override": 0.4,
            "notes": "Заказывать партиями",
        },
    )
    assert create_response.status_code == 200
    item = create_response.json()
    assert item["sku"] == "COF-001"
    assert item["supplier_name"] == "Поставщик SKU"

    list_response = client.get("/sku-settings")
    assert list_response.status_code == 200
    assert len(list_response.json()["items"]) == 1

    single_response = client.get("/sku-settings/COF-001")
    assert single_response.status_code == 200
    assert single_response.json()["moq"] == 30

    delete_response = client.delete("/sku-settings/COF-001")
    assert delete_response.status_code == 200
    missing_response = client.get("/sku-settings/COF-001")
    assert missing_response.status_code == 404


def test_procurement_enrich_plan_applies_supplier_constraints():
    supplier_response = client.post(
        "/suppliers",
        json={"name": "Поставщик Б", "default_lead_time_days": 9},
    )
    supplier_id = supplier_response.json()["id"]

    client.put(
        "/sku-settings/COF-001",
        json={"supplier_id": supplier_id, "purchase_price": 250, "moq": 30, "pack_size": 12},
    )
    client.put(
        "/sku-settings/TEA-001",
        json={"supplier_id": supplier_id, "pack_size": 6},
    )

    response = client.post(
        "/procurement/enrich-plan",
        json={
            "items": [
                {
                    "sku": "COF-001",
                    "product_name": "Кофе",
                    "category": "Кофе",
                    "recommended_order_qty": 17,
                    "purchase_price": 200,
                },
                {
                    "sku": "TEA-001",
                    "product_name": "Чай",
                    "category": "Чай",
                    "recommended_order_qty": 17,
                    "supplier": "Файл поставщика",
                },
                {
                    "sku": "MILK-001",
                    "product_name": "Молоко",
                    "category": "Молоко",
                    "recommended_order_qty": 5,
                },
            ]
        },
    )
    assert response.status_code == 200
    items = {item["sku"]: item for item in response.json()["items"]}

    assert items["COF-001"]["system_recommended_qty"] == 17
    assert items["COF-001"]["adjusted_order_qty"] == 36
    assert "минимального заказа 30" in items["COF-001"]["adjustment_reason"]
    assert "кратности упаковки 12" in items["COF-001"]["adjustment_reason"]
    assert items["COF-001"]["supplier_name"] == "Поставщик Б"
    assert items["COF-001"]["supplier_source"] == "directory"
    assert items["COF-001"]["purchase_price"] == 250
    assert items["COF-001"]["estimated_cost"] == 9000

    assert items["TEA-001"]["adjusted_order_qty"] == 18
    assert items["TEA-001"]["supplier_name"] == "Поставщик Б"
    assert items["TEA-001"]["supplier_lead_time_days"] == 9

    assert items["MILK-001"]["adjusted_order_qty"] == 5
    assert items["MILK-001"]["adjustment_reason"] == "Без дополнительных ограничений"
    assert items["MILK-001"]["supplier_name"] is None
